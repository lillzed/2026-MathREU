"""
500 Card Game - Data Collection & Parser
=========================================
Runs N bot games, captures output, and parses into a normalized,
multi-sheet Excel workbook (Hands / Bids / Tricks / Plays).

Usage:
    python collect_and_parse.py --games 500 --out data/games.xlsx
    python collect_and_parse.py --games 5 --keep-raw
"""

import subprocess
import os
import re
import argparse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


# ─────────────────────────────────────────
# STEP 1: Run games and capture output
# ─────────────────────────────────────────

def run_games(n, server_path, port, password, bot_types, raw_dir=None):
    results = []
    for i in range(n):
        print(f"Running game {i + 1}/{n}...", end="\r")
        result = subprocess.run(
            [server_path, str(port), password, bot_types],
            capture_output=True,
            text=True
        )
        output = result.stdout
        results.append(output)

        if raw_dir:
            os.makedirs(raw_dir, exist_ok=True)
            with open(os.path.join(raw_dir, f"game_{i+1}.txt"), "w") as f:
                f.write(output)

    print(f"\nDone! Ran {n} games.")
    return results


# ─────────────────────────────────────────
# STEP 2: Parse raw output into structured data
# ─────────────────────────────────────────

def parse_output(raw_text):
    """Parse one server run (one or more hands) into a list of hand dicts."""
    hands = []
    lines = raw_text.strip().split("\n")

    current_hand = None
    current_trick = None
    hand_number = 0
    trick_number = 0
    bid_order = 0

    for line in lines:
        line = line.strip()

        if line == "Shuffling and Dealing":
            if current_hand is not None:
                hands.append(current_hand)
            hand_number += 1
            trick_number = 0
            bid_order = 0
            current_hand = {
                "hand_number": hand_number,
                "bids": [],
                "winning_bid": None,
                "kitty": None,
                "discards": [],
                "tricks": [],
                "outcome": None,
                "score_team0": None,
                "score_team1": None,
            }
            current_trick = None

        elif (m := re.match(r"Player (\d+) passed", line)):
            if current_hand:
                bid_order += 1
                current_hand["bids"].append({
                    "order": bid_order, "player": int(m.group(1)),
                    "action": "pass", "bid_value": None
                })

        elif (m := re.match(r"Player (\d+) bet (\S+)", line)):
            if current_hand:
                bid_order += 1
                current_hand["bids"].append({
                    "order": bid_order, "player": int(m.group(1)),
                    "action": "bet", "bid_value": m.group(2)
                })

        elif (m := re.match(r"Player (\d+) won the bet with (\S+)!", line)):
            if current_hand:
                current_hand["winning_bid"] = {"player": int(m.group(1)), "bid": m.group(2)}

        elif (m := re.match(r"Kitty: (\S+) (\S+) (\S+)", line)):
            if current_hand:
                current_hand["kitty"] = [m.group(1), m.group(2), m.group(3)]

        elif (m := re.match(r"Discarded: (\S+)", line)):
            if current_hand:
                current_hand["discards"].append(m.group(1))

        elif (m := re.match(r"Player (\d+) played (\S+)\. Player (\d+) (winning|won) with (\S+)", line)):
            if current_hand is None:
                continue

            player = int(m.group(1))
            card = m.group(2)
            status = m.group(4)
            trick_winner = int(m.group(3))
            winning_card = m.group(5)

            if current_trick is None:
                trick_number += 1
                current_trick = {
                    "trick_number": trick_number,
                    "plays": [],
                    "winner": None,
                    "winning_card": None,
                    "betting_team_tricks": None,
                }

            current_trick["plays"].append({"player": player, "card": card})

            if status == "won":
                current_trick["winner"] = trick_winner
                current_trick["winning_card"] = winning_card
                current_hand["tricks"].append(current_trick)
                current_trick = None

        elif (m := re.match(r"Betting team has won (\d+) trick\(s\)", line)):
            if current_hand and current_hand["tricks"]:
                current_hand["tricks"][-1]["betting_team_tricks"] = int(m.group(1))

        elif line == "Betting team won!":
            if current_hand:
                current_hand["outcome"] = "won"

        elif line == "Betting team lost!":
            if current_hand:
                current_hand["outcome"] = "lost"

        elif (m := re.match(r"Team 0 points: (-?\d+)", line)):
            if current_hand:
                current_hand["score_team0"] = int(m.group(1))

        elif (m := re.match(r"Team 1 points: (-?\d+)", line)):
            if current_hand:
                current_hand["score_team1"] = int(m.group(1))

    if current_hand is not None:
        hands.append(current_hand)

    return hands


# ─────────────────────────────────────────
# STEP 3: Build normalized row sets
# ─────────────────────────────────────────

def build_tables(all_games_hands):
    """
    all_games_hands: list of (game_number, list_of_hand_dicts)
    Returns dict of table_name -> list of rows (each row a list matching that table's headers)
    """
    hands_rows = []
    bids_rows = []
    tricks_rows = []
    plays_rows = []

    for game_number, hands in all_games_hands:
        for hand in hands:
            winning_bid = hand.get("winning_bid") or {}
            bid_winner = winning_bid.get("player")
            bid_value = winning_bid.get("bid")
            kitty = hand.get("kitty") or [None, None, None]
            discards = hand.get("discards") or [None, None, None]
            while len(discards) < 3:
                discards.append(None)

            # Hands table — one row per hand
            hands_rows.append([
                game_number,
                hand["hand_number"],
                bid_winner,
                bid_winner % 2 if bid_winner is not None else None,
                bid_value,
                kitty[0], kitty[1], kitty[2],
                discards[0], discards[1], discards[2],
                hand["outcome"],
                hand["score_team0"],
                hand["score_team1"],
            ])

            # Bids table — one row per bid/pass action
            for bid in hand["bids"]:
                bids_rows.append([
                    game_number,
                    hand["hand_number"],
                    bid["order"],
                    bid["player"],
                    bid["player"] % 2,
                    bid["action"],
                    bid["bid_value"],
                ])

            # Tricks table — one row per trick
            for trick in hand["tricks"]:
                tricks_rows.append([
                    game_number,
                    hand["hand_number"],
                    trick["trick_number"],
                    trick["winner"],
                    trick["winning_card"],
                    trick["betting_team_tricks"],
                ])

                # Plays table — one row per card played
                for play_index, play in enumerate(trick["plays"]):
                    plays_rows.append([
                        game_number,
                        hand["hand_number"],
                        trick["trick_number"],
                        play_index,
                        play["player"],
                        play["player"] % 2,
                        play["card"],
                        1 if play["player"] == trick["winner"] else 0,
                    ])

    return {
        "Hands": hands_rows,
        "Bids": bids_rows,
        "Tricks": tricks_rows,
        "Plays": plays_rows,
    }


# ─────────────────────────────────────────
# STEP 4: Write to Excel (multi-sheet)
# ─────────────────────────────────────────

HEADERS = {
    "Hands": [
        "game_number", "hand_number", "bid_winner", "bid_winner_team", "bid_value",
        "kitty_card_1", "kitty_card_2", "kitty_card_3",
        "discard_1", "discard_2", "discard_3",
        "outcome", "final_score_team0", "final_score_team1",
    ],
    "Bids": [
        "game_number", "hand_number", "bid_order", "player", "player_team",
        "action", "bid_value",
    ],
    "Tricks": [
        "game_number", "hand_number", "trick_number",
        "winner", "winning_card", "betting_team_tricks_after",
    ],
    "Plays": [
        "game_number", "hand_number", "trick_number", "play_index",
        "player", "player_team", "card_played", "won_trick",
    ],
}


def write_sheet(wb, name, headers, rows):
    ws = wb.create_sheet(name) if name in wb.sheetnames and wb[name].max_row == 1 and wb[name]["A1"].value is None else wb.create_sheet(name)
    ws = wb[name]

    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", start_color="2E4057")
    header_align = Alignment(horizontal="center")

    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    row_font = Font(name="Arial", size=10)
    for row_data in rows:
        ws.append(row_data)

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.font = row_font

    for col in ws.columns:
        max_len = max((len(str(c.value)) if c.value is not None else 0) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 20)

    ws.freeze_panes = "A2"


def write_excel(tables, out_path):
    wb = Workbook()
    wb.remove(wb.active)  # remove default empty sheet

    for name in ["Hands", "Bids", "Tricks", "Plays"]:
        write_sheet(wb, name, HEADERS[name], tables[name])

    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    wb.save(out_path)

    for name in ["Hands", "Bids", "Tricks", "Plays"]:
        print(f"  {name}: {len(tables[name])} rows")
    print(f"Saved to {out_path}")


# ─────────────────────────────────────────
# MAIN
# ─────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(description="Collect and parse 500 card game data")
    parser.add_argument("--games", type=int, default=100, help="Number of games to run")
    parser.add_argument("--out", type=str, default="data/games.xlsx", help="Output Excel file")
    parser.add_argument("--server", type=str, default="./500/server", help="Path to server binary")
    parser.add_argument("--port", type=int, default=8080, help="Port for server")
    parser.add_argument("--password", type=str, default="pass", help="Server password")
    parser.add_argument("--bots", type=str, default="2222", help="Bot config e.g. 2222")
    parser.add_argument("--keep-raw", action="store_true", help="Save raw .txt files to data/raw/")
    args = parser.parse_args()

    raw_dir = "data/raw" if args.keep_raw else None

    raw_outputs = run_games(
        n=args.games,
        server_path=args.server,
        port=args.port,
        password=args.password,
        bot_types=args.bots,
        raw_dir=raw_dir
    )

    all_games_hands = []
    for game_number, raw in enumerate(raw_outputs, start=1):
        hands = parse_output(raw)
        all_games_hands.append((game_number, hands))

    tables = build_tables(all_games_hands)
    write_excel(tables, args.out)


if __name__ == "__main__":
    main()